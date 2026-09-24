#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
QCS Labo パラメータシート生成スクリプト
=============================================================
group_vars/all.yml を読み込み、設定値の定義書(Excel)を生成します。
値は all.yml を唯一のソースとするため、Ansible の適用値と
Excel パラメータシートが常に一致します。

前提:
    pip install pyyaml openpyxl

使い方（このPC / 踏み台どちらでも）:
    python generate_param_sheet.py
    # 出力: docs/QCS_Labo_パラメータシート.xlsx
"""

import os
import sys
import datetime

try:
    import yaml
except ImportError:
    sys.exit("pyyaml が必要です: pip install pyyaml")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl が必要です: pip install openpyxl")


# -------------------------------------------------------------
# パス設定（このスクリプトの位置を基準に all.yml を探す）
# -------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
ALL_YML = os.path.normpath(os.path.join(HERE, "..", "ansible", "group_vars", "all.yml"))
OUT_XLSX = os.path.join(HERE, "QCS_Labo_パラメータシート.xlsx")


# -------------------------------------------------------------
# スタイル定義
# -------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="305496")   # 濃紺
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")   # 薄青
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def load_vars(path):
    """all.yml を読み込んで dict で返す。"""
    if not os.path.exists(path):
        sys.exit(f"変数ファイルが見つかりません: {path}")
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def style_header_row(ws, row, ncols):
    """指定行をヘッダー体裁にする。"""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def put_row(ws, row, values, border=True):
    """1 行分の値を書き込む。"""
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.alignment = WRAP
        if border:
            cell.border = BORDER


def section_title(ws, row, text, ncols):
    """セクション見出し行。"""
    ws.cell(row=row, column=1, value=text)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
        cell.border = BORDER
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)


def autofit(ws, widths):
    """列幅を設定。"""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# =============================================================
# シート1: 概要
# =============================================================
def build_overview(wb, data):
    ws = wb.active
    ws.title = "概要"
    meta = data.get("project_meta", {})

    ws["A1"] = meta.get("project_name", "QCS Labo パラメータシート")
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    rows = [
        ("項目", "値"),
        ("プロジェクト名", meta.get("project_name", "")),
        ("クラウド基盤", meta.get("platform", "")),
        ("対象 OS", meta.get("os", "")),
        ("作成者", meta.get("author", "")),
        ("生成日時", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("設定ソース", "ansible/group_vars/all.yml"),
    ]
    for r, (k, v) in enumerate(rows, start=3):
        put_row(ws, r, [k, v])
    style_header_row(ws, 3, 2)
    autofit(ws, [22, 55])


# =============================================================
# シート2: ユーザー / sudo
# =============================================================
def build_users(wb, data):
    ws = wb.create_sheet("ユーザー・sudo")
    users = data.get("users", [])
    policies = data.get("sudo_policies", {})
    ncols = 6

    section_title(ws, 1, "ユーザー登録一覧", ncols)
    put_row(ws, 2, ["ユーザー名", "説明", "所属グループ", "sudo", "sudo方針", "ログインシェル"])
    style_header_row(ws, 2, ncols)

    policy_label = {
        "A": "方針A（ブロックリスト方式）",
        "B": "方針B（ホワイトリスト方式）",
        "unrestricted": "制限なし（フルsudo）",
    }
    r = 3
    for u in users:
        put_row(ws, r, [
            u.get("name", ""),
            u.get("comment", ""),
            ", ".join(u.get("groups", []) or []),
            "有効" if u.get("sudo") else "無効",
            policy_label.get(u.get("sudo_policy", ""), u.get("sudo_policy", "")),
            u.get("shell", ""),
        ])
        r += 1

    # sudo 方針の詳細
    r += 1
    section_title(ws, r, "sudo 方針の定義", ncols)
    r += 1
    put_row(ws, r, ["方針", "種別", "対象コマンド", "", "", ""])
    style_header_row(ws, r, ncols)
    r += 1

    a = policies.get("A", {})
    put_row(ws, r, [
        "方針A", "禁止コマンド（denied）",
        "\n".join(a.get("denied_commands", []) or []), "", "", "",
    ])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ncols)
    r += 1

    b = policies.get("B", {})
    put_row(ws, r, [
        "方針B", "許可コマンド（allowed）",
        "\n".join(b.get("allowed_commands", []) or []), "", "", "",
    ])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ncols)

    autofit(ws, [16, 34, 22, 12, 26, 16])


# =============================================================
# シート3: システム設定（SELinux / パッケージ / OS最新化）
# =============================================================
def build_system(wb, data):
    ws = wb.create_sheet("システム設定")
    ncols = 3

    section_title(ws, 1, "SELinux", ncols)
    put_row(ws, 2, ["項目", "値", "備考"])
    style_header_row(ws, 2, ncols)
    put_row(ws, 3, ["SELinux 状態", data.get("selinux_state", ""),
                    "disabled は再起動後に完全反映"])

    r = 5
    section_title(ws, r, "導入パッケージ", ncols)
    r += 1
    put_row(ws, r, ["項目", "値", "備考"])
    style_header_row(ws, r, ncols)
    r += 1
    put_row(ws, r, ["EPEL 有効化", "有効" if data.get("enable_epel") else "無効",
                    "stress-ng の導入に必要"])
    r += 1
    pkg_notes = {
        "iperf3": "ネットワーク帯域測定",
        "fio": "ストレージ I/O ベンチマーク",
        "stress-ng": "CPU/メモリ等の負荷試験（EPEL）",
    }
    for pkg in data.get("packages", []):
        put_row(ws, r, ["パッケージ", pkg, pkg_notes.get(pkg, "")])
        r += 1

    r += 1
    section_title(ws, r, "OS 最新化", ncols)
    r += 1
    put_row(ws, r, ["項目", "値", "備考"])
    style_header_row(ws, r, ncols)
    r += 1
    put_row(ws, r, ["OS 最新化", "実施" if data.get("os_update") else "しない",
                    "dnf update -*"])
    r += 1
    put_row(ws, r, ["要再起動時の自動再起動",
                    "する" if data.get("reboot_if_required") else "しない",
                    "needs-restarting で判定"])

    autofit(ws, [24, 26, 34])


def main():
    data = load_vars(ALL_YML)
    wb = Workbook()
    build_overview(wb, data)
    build_users(wb, data)
    build_system(wb, data)
    wb.save(OUT_XLSX)
    print(f"パラメータシートを生成しました: {OUT_XLSX}")


if __name__ == "__main__":
    main()
